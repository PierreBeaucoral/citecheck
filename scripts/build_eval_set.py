#!/usr/bin/env python3
"""Build data/eval/labels.csv from the Walters & Wilder (2023) supplementary data.

The Walters & Wilder paper published 636 ChatGPT-generated citations with
explicit Fabricated/Real labels (Sci Reports 13:14045, CC-BY). We use this as
the methodologically-clean anchor for the citecheck hallucination eval —
third-party-generated, third-party-labeled.

Usage:
    docker compose up -d grobid                    # GROBID must be reachable
    uv run python scripts/build_eval_set.py        # writes data/eval/labels.csv

Each XLSX row becomes one CSV row. We parse the raw citation string through
GROBID to get title/authors/year/journal where possible — citecheck's
hallucination detector needs those fields. Failures fall back to raw_text only.
"""

from __future__ import annotations

import csv
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

# Add src/ for local-package imports without needing `uv pip install -e .` to be current.
REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))

from citecheck.extraction.grobid_client import parse_citation_string  # noqa: E402

SUPP_XLSX = REPO_ROOT / "data" / "eval" / "raw" / "walters_wilder_2023_supp3.xlsx"
LABELS_CSV = REPO_ROOT / "data" / "eval" / "labels.csv"
SHEET = "Works cited in those papers"

# Column indices (1-based) in the supplement.
COL_GPT = 1
COL_FIELD = 2
COL_TOPIC = 3
COL_CITN = 4
COL_TEXT = 5
COL_SCHOLARLY = 8
COL_TYPE = 9
COL_FABRICATED = 14
COL_NOTES = 22


def _expected_verdict(is_fabricated: int) -> str:
    # Map Walters' boolean "Work itself is fabricated" to citecheck's
    # ground-truth verdict. We use likely_hallucinated for fabricated, and
    # collapse the two "real" tiers to real_high_confidence — citecheck's
    # actual verdict may land on real_low (no DOI in citation), which we
    # accept as a correct prediction during eval scoring.
    return "likely_hallucinated" if is_fabricated == 1 else "real_high_confidence"


def main() -> int:
    if not SUPP_XLSX.is_file():
        print(f"missing: {SUPP_XLSX}", file=sys.stderr)
        return 2

    wb = load_workbook(SUPP_XLSX, data_only=True)
    ws = wb[SHEET]

    rows_written = 0
    fabricated = 0
    real = 0
    parse_failures = 0

    LABELS_CSV.parent.mkdir(parents=True, exist_ok=True)
    with LABELS_CSV.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "id",
                "raw_text",
                "class",
                "ground_truth_doi",
                "expected_verdict",
                "arxiv_id",
                "biorxiv_id",
                "notes",
                "labeler",
                "verified_by",
                "parsed_title",
                "parsed_year",
                "parsed_journal",
                "parsed_authors",
                "parsed_doi",
            ]
        )
        for r in range(2, ws.max_row + 1):
            citation = ws.cell(row=r, column=COL_TEXT).value
            if not citation or not isinstance(citation, str) or not citation.strip():
                continue
            fab_val = ws.cell(row=r, column=COL_FABRICATED).value
            try:
                is_fab = int(fab_val)
            except (TypeError, ValueError):
                continue
            gpt = ws.cell(row=r, column=COL_GPT).value
            field = ws.cell(row=r, column=COL_FIELD).value
            topic = ws.cell(row=r, column=COL_TOPIC).value
            citn = ws.cell(row=r, column=COL_CITN).value
            ptype = ws.cell(row=r, column=COL_TYPE).value
            row_notes = ws.cell(row=r, column=COL_NOTES).value or ""

            try:
                parsed = parse_citation_string(citation.strip())
            except Exception as exc:
                print(f"  parse failure row {r}: {exc}", file=sys.stderr)
                parse_failures += 1
                from citecheck.models import RawReference

                parsed = RawReference(raw_text=citation.strip())

            authors_str = "; ".join(f"{a.family},{a.given or ''}" for a in parsed.authors)
            writer.writerow(
                [
                    f"walters_{r - 1:04d}",
                    citation.strip(),
                    "hallucinated" if is_fab else "real_external",
                    "",  # ground_truth_doi — Walters did not provide DOIs systematically
                    _expected_verdict(is_fab),
                    "",  # arxiv_id — not in source
                    "",  # biorxiv_id
                    f"walters_wilder_2023; gpt={gpt}; field={field}; topic={topic}; "
                    f"citn={citn}; type={ptype}; {row_notes}".strip("; "),
                    "walters_wilder_2023",
                    "walters_wilder_2023",
                    parsed.title or "",
                    parsed.year or "",
                    parsed.journal or "",
                    authors_str,
                    parsed.doi or "",
                ]
            )
            rows_written += 1
            if is_fab:
                fabricated += 1
            else:
                real += 1
            # Be polite to GROBID — micro-throttle to avoid runaway concurrency.
            if rows_written % 50 == 0:
                print(f"  ... {rows_written} rows written", file=sys.stderr)
                time.sleep(0.05)

    print(
        f"wrote {rows_written} rows to {LABELS_CSV.relative_to(REPO_ROOT)}\n"
        f"  real: {real}\n"
        f"  fabricated: {fabricated}\n"
        f"  citation parse failures (fell back to raw_text only): {parse_failures}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
